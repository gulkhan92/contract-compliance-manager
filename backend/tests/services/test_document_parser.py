import io

import docx
import pymupdf

from app.services.document_parser import parse_docx, parse_pdf


def _build_pdf(lines: list[str]) -> bytes:
    # PyMuPDF ships no type stubs.
    document = pymupdf.open()  # type: ignore[no-untyped-call]
    page = document.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + i * 30), line)
    data: bytes = document.tobytes()  # type: ignore[no-untyped-call]
    document.close()  # type: ignore[no-untyped-call]
    return data


def _build_docx(lines: list[str], *, heading_index: int = 0) -> bytes:
    document = docx.Document()
    for i, line in enumerate(lines):
        if i == heading_index:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_LINES = [
    "TERM AND RENEWAL",
    "This Agreement renews automatically for successive 1 year terms unless "
    "terminated with 90 days notice.",
    "IN WITNESS WHEREOF, the parties have executed this Agreement.",
]


def test_parse_pdf_extracts_paragraphs_in_order() -> None:
    paragraphs = parse_pdf(_build_pdf(_LINES))

    assert [p.raw_text for p in paragraphs] == _LINES
    assert [p.paragraph_index for p in paragraphs] == [0, 1, 2]


def test_parse_pdf_tags_heading_onto_following_paragraphs() -> None:
    paragraphs = parse_pdf(_build_pdf(_LINES))

    assert all(p.section_heading == "TERM AND RENEWAL" for p in paragraphs)


def test_parse_docx_extracts_paragraphs_in_order() -> None:
    paragraphs = parse_docx(_build_docx(_LINES))

    assert [p.raw_text for p in paragraphs] == _LINES


def test_parse_docx_uses_heading_style_not_just_text_shape() -> None:
    lines = ["Term and Renewal", "This Agreement renews for one additional year."]
    paragraphs = parse_docx(_build_docx(lines, heading_index=0))

    # Mixed-case text wouldn't trip the ALL-CAPS heuristic — the docx
    # "Heading 1" style itself is what should be picked up here.
    assert paragraphs[0].section_heading == "Term and Renewal"
    assert paragraphs[1].section_heading == "Term and Renewal"


def test_parse_docx_skips_empty_paragraphs() -> None:
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph.")
    buffer = io.BytesIO()
    document.save(buffer)

    paragraphs = parse_docx(buffer.getvalue())

    assert [p.raw_text for p in paragraphs] == ["First paragraph.", "Second paragraph."]
