import io
import uuid

import docx
import pymupdf
import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import create_access_token
from app.db.enums import UserRole
from app.db.models import ContractChunk, User
from app.services.file_validation import MAX_UPLOAD_SIZE_BYTES


def _build_pdf_bytes(lines: list[str]) -> bytes:
    document = pymupdf.open()  # type: ignore[no-untyped-call]
    page = document.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + i * 30), line)
    data: bytes = document.tobytes()  # type: ignore[no-untyped-call]
    document.close()  # type: ignore[no-untyped-call]
    return data


_RENEWAL_CONTRACT_PDF = _build_pdf_bytes(
    [
        "TERM AND RENEWAL",
        "This Agreement renews automatically for successive 1 year terms unless",
        "either party gives 90 days written notice of non-renewal.",
    ]
)


async def _register_and_login(client: AsyncClient, *, org_name: str, email: str) -> str:
    response = await client.post(
        "/api/v1/auth/register",
        json={
            "org_name": org_name,
            "email": email,
            "password": "correct horse battery staple",
            "full_name": "Test Admin",
        },
    )
    assert response.status_code == 201, response.text
    token: str = response.json()["access_token"]
    return token


def _auth_headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.asyncio
async def test_upload_requires_editor_role(client: AsyncClient, db_session: AsyncSession) -> None:
    admin_token = await _register_and_login(client, org_name="Acme", email="admin1@example.com")
    me = (await client.get("/api/v1/auth/me", headers=_auth_headers(admin_token))).json()

    viewer = User(
        org_id=uuid.UUID(me["org_id"]),
        email="viewer1@example.com",
        hashed_password="irrelevant",
        role=UserRole.VIEWER,
        full_name="A Viewer",
    )
    db_session.add(viewer)
    await db_session.flush()


    viewer_token = create_access_token(user_id=viewer.id, org_id=viewer.org_id, role=viewer.role)

    response = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(viewer_token),
        files={"file": ("contract.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")},
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_upload_creates_contract_with_chunks_and_prefilter_hits(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin2@example.com")

    response = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={"file": ("renewal.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")},
        data={"title": "Renewal Agreement"},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["contract"]["title"] == "Renewal Agreement"
    assert body["contract"]["status"] == "processing"
    assert body["extraction_job_id"]

    contract_id = body["contract"]["id"]
    chunks = (
        await db_session.execute(
            select(ContractChunk).where(ContractChunk.contract_id == contract_id)
        )
    ).scalars().all()
    assert len(chunks) == 3
    assert any(c.passed_prefilter for c in chunks)
    assert any(c.section_heading == "TERM AND RENEWAL" for c in chunks)


@pytest.mark.asyncio
async def test_upload_rejects_oversized_file(client: AsyncClient) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin3@example.com")

    oversized = b"%PDF-1.7\n" + b"x" * MAX_UPLOAD_SIZE_BYTES

    response = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={"file": ("big.pdf", oversized, "application/pdf")},
    )
    assert response.status_code == 413


@pytest.mark.asyncio
async def test_upload_rejects_invalid_file_type(client: AsyncClient) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin4@example.com")

    response = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={"file": ("fake.pdf", b"not actually a pdf", "application/pdf")},
    )
    assert response.status_code == 415


@pytest.mark.asyncio
async def test_upload_duplicate_file_conflicts(client: AsyncClient) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin5@example.com")
    headers = _auth_headers(token)
    files = {"file": ("renewal.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")}

    first = await client.post("/api/v1/contracts", headers=headers, files=files)
    assert first.status_code == 201

    second = await client.post("/api/v1/contracts", headers=headers, files=files)
    assert second.status_code == 409


@pytest.mark.asyncio
async def test_contracts_are_scoped_to_org(client: AsyncClient) -> None:
    """The plan's explicit cross-org isolation acceptance test, now that a
    real org-scoped resource (contracts) exists — see
    docs/CONTRACT_CLM_BUILD_PLAN.md §12 Phase 2's test list."""
    token_a = await _register_and_login(client, org_name="Org A", email="a-admin@example.com")
    token_b = await _register_and_login(client, org_name="Org B", email="b-admin@example.com")

    upload = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token_a),
        files={"file": ("renewal.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")},
    )
    assert upload.status_code == 201
    contract_id = upload.json()["contract"]["id"]

    # Org B must not see org A's contract in its list...
    list_response = await client.get("/api/v1/contracts", headers=_auth_headers(token_b))
    assert list_response.json() == []

    # ...nor be able to fetch it directly by id.
    get_response = await client.get(
        f"/api/v1/contracts/{contract_id}", headers=_auth_headers(token_b)
    )
    assert get_response.status_code == 404

    # Org A can see its own contract.
    own_list = await client.get("/api/v1/contracts", headers=_auth_headers(token_a))
    assert len(own_list.json()) == 1


@pytest.mark.asyncio
async def test_get_contract_status_reports_latest_extraction_job(client: AsyncClient) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin6@example.com")

    upload = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={"file": ("renewal.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")},
    )
    contract_id = upload.json()["contract"]["id"]

    response = await client.get(
        f"/api/v1/contracts/{contract_id}/status", headers=_auth_headers(token)
    )
    assert response.status_code == 200
    body = response.json()
    assert body["contract_status"] == "processing"
    assert body["latest_extraction_job"]["status"] == "succeeded"


@pytest.mark.asyncio
async def test_delete_contract_requires_editor_role(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin7@example.com")
    me = (await client.get("/api/v1/auth/me", headers=_auth_headers(token))).json()

    upload = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={"file": ("renewal.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")},
    )
    contract_id = upload.json()["contract"]["id"]

    viewer = User(
        org_id=uuid.UUID(me["org_id"]),
        email="viewer2@example.com",
        hashed_password="irrelevant",
        role=UserRole.VIEWER,
        full_name="A Viewer",
    )
    db_session.add(viewer)
    await db_session.flush()


    viewer_token = create_access_token(user_id=viewer.id, org_id=viewer.org_id, role=viewer.role)

    response = await client.delete(
        f"/api/v1/contracts/{contract_id}", headers=_auth_headers(viewer_token)
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_delete_contract_cascades_chunks(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin8@example.com")

    upload = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={"file": ("renewal.pdf", _RENEWAL_CONTRACT_PDF, "application/pdf")},
    )
    contract_id = upload.json()["contract"]["id"]

    response = await client.delete(
        f"/api/v1/contracts/{contract_id}", headers=_auth_headers(token)
    )
    assert response.status_code == 204

    remaining_chunks = (
        await db_session.execute(
            select(ContractChunk).where(ContractChunk.contract_id == contract_id)
        )
    ).scalars().all()
    assert remaining_chunks == []


@pytest.mark.asyncio
async def test_upload_accepts_docx(client: AsyncClient) -> None:
    token = await _register_and_login(client, org_name="Acme", email="admin9@example.com")

    document = docx.Document()
    document.add_heading("Governing Law", level=1)
    document.add_paragraph("This Agreement is governed by the laws of the State of New York.")
    buffer = io.BytesIO()
    document.save(buffer)

    response = await client.post(
        "/api/v1/contracts",
        headers=_auth_headers(token),
        files={
            "file": (
                "contract.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
